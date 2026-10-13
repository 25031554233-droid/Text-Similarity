import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
import seaborn as sns
from sklearn.feature_extraction.text import CountVectorizer
from scipy.spatial.distance import pdist, squareform
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.drawing.image import Image

# ---------------------------------------------------------
# 1. CLASS ENGINE: ANALISIS TEKS & METRIK
# ---------------------------------------------------------
class DocumentSimilarityEngine:
    def __init__(self, corpus):
        self.corpus = corpus
        self.doc_ids = list(corpus.keys())
        self.texts = list(corpus.values())
        
        # Vectorizer menggunakan Scikit-Learn
        self.vectorizer = CountVectorizer()
        self.dtm = self.vectorizer.fit_transform(self.texts).toarray()
        self.feature_names = self.vectorizer.get_feature_names_out()
        
        self.df_dtm = pd.DataFrame(self.dtm, index=self.doc_ids, columns=self.feature_names)
        self.matrices = {}
        
    def compute_similarity_and_distances(self):
        # Cosine Similarity (Vectorized Dot Product)
        norms = np.linalg.norm(self.dtm, axis=1, keepdims=True)
        norms[norms == 0] = 1.0
        normalized_dtm = self.dtm / norms
        cosine_sim = np.dot(normalized_dtm, normalized_dtm.T)
        self.matrices['cosine'] = pd.DataFrame(cosine_sim, index=self.doc_ids, columns=self.doc_ids)

        # Jaccard Similarity (Intersection over Union)
        n = len(self.texts)
        jaccard_sim = np.zeros((n, n))
        sets = [set(t.split()) for t in self.texts]
        for i in range(n):
            for j in range(n):
                inter = len(sets[i] & sets[j])
                union = len(sets[i] | sets[j])
                jaccard_sim[i, j] = inter / union if union > 0 else 0.0
        self.matrices['jaccard'] = pd.DataFrame(jaccard_sim, index=self.doc_ids, columns=self.doc_ids)

        # Distance Metrics via SciPy (Euclidean, Manhattan, Chebyshev)
        for metric_name, scipy_metric in [
            ('euclidean', 'euclidean'),
            ('manhattan', 'cityblock'),
            ('chebyshev', 'chebyshev')
        ]:
            dist_array = squareform(pdist(self.dtm, metric=scipy_metric))
            self.matrices[metric_name] = pd.DataFrame(dist_array, index=self.doc_ids, columns=self.doc_ids)

        # Minkowski Distance (p=3)
        minkowski_dist = squareform(pdist(self.dtm, metric='minkowski', p=3))
        self.matrices['minkowski'] = pd.DataFrame(minkowski_dist, index=self.doc_ids, columns=self.doc_ids)

    def generate_heatmap(self, output_img="heatmap_similarity_matrix.png"):
        fig, axes = plt.subplots(1, 2, figsize=(14, 5.8))
        sns.set_theme(style="white")

        # Plot Cosine
        sns.heatmap(self.matrices['cosine'], annot=True, fmt=".3f", cmap="YlGnBu", ax=axes[0],
                    cbar=True, linewidths=0.5, annot_kws={"size": 9.5, "weight": "bold"})
        axes[0].set_title("Cosine Similarity Matrix", fontsize=12, fontweight="bold", pad=10)

        # Plot Jaccard
        sns.heatmap(self.matrices['jaccard'], annot=True, fmt=".3f", cmap="Purples", ax=axes[1],
                    cbar=True, linewidths=0.5, annot_kws={"size": 9.5, "weight": "bold"})
        axes[1].set_title("Jaccard Similarity Matrix", fontsize=12, fontweight="bold", pad=10)

        plt.tight_layout()
        plt.savefig(output_img, dpi=300, bbox_inches='tight')
        plt.close()
        return output_img


# ---------------------------------------------------------
# 2. CLASS EXPORTER: DASHBOARD & GENERATOR EXCEL
# ---------------------------------------------------------
class ExcelDashboardExporter:
    def __init__(self, engine, image_path):
        self.engine = engine
        self.image_path = image_path
        self.wb = openpyxl.Workbook()
        
        # Styling Palette Modern
        self.style_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        self.style_title = Font(name="Calibri", size=16, bold=True, color="1E293B")
        self.style_sec = Font(name="Calibri", size=12, bold=True, color="334155")
        self.font_data = Font(name="Calibri", size=10)
        
        self.fill_primary = PatternFill("solid", fgColor="1E293B")   # Charcoal
        self.fill_secondary = PatternFill("solid", fgColor="4338CA") # Indigo
        self.fill_tertiary = PatternFill("solid", fgColor="047857")  # Emerald
        self.fill_kpi = PatternFill("solid", fgColor="F8FAFC")
        self.fill_alt = PatternFill("solid", fgColor="F1F5F9")

        self.thin_border = Border(
            left=Side(style='thin', color='CBD5E1'), right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'), bottom=Side(style='thin', color='CBD5E1')
        )
        self.kpi_border = Border(
            left=Side(style='medium', color='94A3B8'), right=Side(style='medium', color='94A3B8'),
            top=Side(style='medium', color='94A3B8'), bottom=Side(style='medium', color='94A3B8')
        )

    def _apply_format(self, ws, s_row, s_col, e_row, e_col, fill, num_fmt=None, align="center"):
        for r in range(s_row, e_row + 1):
            for c in range(s_col, e_col + 1):
                cell = ws.cell(row=r, column=c)
                cell.border = self.thin_border
                if r == s_row:
                    cell.font = self.style_header
                    cell.fill = fill
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.font = self.font_data
                    if r % 2 == 1:
                        cell.fill = self.fill_alt
                    cell.alignment = Alignment(horizontal=align, vertical="center")
                    if num_fmt and c > s_col:
                        cell.number_format = num_fmt

    def build_overview_sheet(self):
        ws = self.wb.active
        ws.title = "Overview & Summary"
        ws.sheet_properties.tabColor = "1E293B"
        ws.views.sheetView[0].showGridLines = True

        # Judul Utama
        ws.cell(row=2, column=2, value="TEXT SIMILARITY ANALYTICS DASHBOARD").font = self.style_title
        ws.cell(row=3, column=2, value="Systematic Evaluation of Document Distances and Semantic Overlap").font = Font(name="Calibri", size=10, italic=True, color="64748B")

        # KPI Cards
        kpi_specs = [
            ("TOTAL DOKUMEN", str(len(self.engine.doc_ids)), 2),
            ("VOCABULARY SIZE", str(len(self.engine.feature_names)), 5),
            ("AVG COSINE SIMILARITY", f"{self.engine.matrices['cosine'].values[np.triu_indices(len(self.engine.doc_ids), k=1)].mean():.3f}", 8)
        ]

        for label, val, start_col in kpi_specs:
            ws.merge_cells(start_row=5, start_column=start_col, end_row=5, end_column=start_col + 1)
            ws.merge_cells(start_row=6, start_column=start_col, end_row=6, end_column=start_col + 1)
            
            c_lbl = ws.cell(row=5, column=start_col, value=label)
            c_lbl.font = Font(name="Calibri", size=9, bold=True, color="64748B")
            c_lbl.fill = self.fill_kpi
            c_lbl.alignment = Alignment(horizontal="center", vertical="center")
            
            c_val = ws.cell(row=6, column=start_col, value=val)
            c_val.font = Font(name="Calibri", size=18, bold=True, color="1E293B")
            c_val.fill = self.fill_kpi
            c_val.alignment = Alignment(horizontal="center", vertical="center")

            for r in range(5, 7):
                for c in range(start_col, start_col + 2):
                    ws.cell(row=r, column=c).border = self.kpi_border

        # Tabel Korpus Dokumen
        ws.cell(row=8, column=2, value="1. Document Corpus Overview").font = self.style_sec
        ws.cell(row=9, column=2, value="ID")
        ws.cell(row=9, column=3, value="Document Text Content")
        ws.cell(row=9, column=4, value="Word Count")

        curr_r = 10
        for doc_id, text in self.engine.corpus.items():
            ws.cell(row=curr_r, column=2, value=doc_id)
            ws.cell(row=curr_r, column=3, value=text)
            ws.cell(row=curr_r, column=4, value=len(text.split()))
            curr_r += 1
        self._apply_format(ws, 9, 2, curr_r - 1, 4, self.fill_primary, align="left")

        # Tabel Benchmark Pasangan Dokumen
        ws.cell(row=curr_r + 1, column=2, value="2. Highlighted Pairwise Benchmarks").font = self.style_sec
        headers = ["Document Pair", "Cosine Score", "Jaccard Score", "Euclidean Dist.", "Interpretation"]
        for c_i, h in enumerate(headers, start=2):
            ws.cell(row=curr_r + 2, column=c_i, value=h)

        pairs = [
            ("D2 vs D5", "cosine", "jaccard", "euclidean", "High Semantic Overlap (AI & NLP)"),
            ("D1 vs D4", "cosine", "jaccard", "euclidean", "Strong Machine Learning Context"),
            ("D3 vs D8", "cosine", "jaccard", "euclidean", "Moderate Match (Cosine / Mining)"),
            ("D6 vs D7", "cosine", "jaccard", "euclidean", "Low Correlation (Distinct Fields)")
        ]

        p_row = curr_r + 3
        for pair_name, m1, m2, m3, desc in pairs:
            d1, d2 = pair_name.split(" vs ")
            val1 = self.engine.matrices[m1].loc[d1, d2]
            val2 = self.engine.matrices[m2].loc[d1, d2]
            val3 = self.engine.matrices[m3].loc[d1, d2]
            
            ws.cell(row=p_row, column=2, value=pair_name)
            ws.cell(row=p_row, column=3, value=val1)
            ws.cell(row=p_row, column=4, value=val2)
            ws.cell(row=p_row, column=5, value=val3)
            ws.cell(row=p_row, column=6, value=desc)
            p_row += 1

        self._apply_format(ws, curr_r + 2, 2, p_row - 1, 6, self.fill_primary, num_fmt="0.0000", align="left")

        # Embed Gambar Heatmap
        ws.cell(row=p_row + 1, column=2, value="3. Heatmap Matrix Visualization").font = self.style_sec
        img = Image(self.image_path)
        img.width = 700
        img.height = 290
        ws.add_image(img, f"B{p_row + 3}")

    def build_matrices_sheet(self):
        ws = self.wb.create_sheet(title="Similarity Matrices")
        ws.sheet_properties.tabColor = "4338CA"
        ws.views.sheetView[0].showGridLines = True

        def _write_matrix(start_row, title, key, fill):
            ws.cell(row=start_row, column=2, value=title).font = self.style_sec
            ws.cell(row=start_row + 1, column=2, value="Docs")
            for c_i, d_id in enumerate(self.engine.doc_ids, start=3):
                ws.cell(row=start_row + 1, column=c_i, value=d_id)
            
            r_idx = start_row + 2
            df_mat = self.engine.matrices[key]
            for d1 in self.engine.doc_ids:
                ws.cell(row=r_idx, column=2, value=d1)
                for c_i, d2 in enumerate(self.engine.doc_ids, start=3):
                    ws.cell(row=r_idx, column=c_i, value=df_mat.loc[d1, d2])
                r_idx += 1
            
            self._apply_format(ws, start_row + 1, 2, r_idx - 1, len(self.engine.doc_ids) + 2, fill, num_fmt="0.0000")
            return r_idx + 2

        r = _write_matrix(2, "1. COSINE SIMILARITY MATRIX", 'cosine', self.fill_secondary)
        cs_blue = ColorScaleRule(start_type='num', start_value=0, start_color='FFFFFF',
                                 end_type='num', end_value=1, end_color='818CF8')
        ws.conditional_formatting.add(f"C4:J11", cs_blue)

        r = _write_matrix(r, "2. JACCARD SIMILARITY MATRIX", 'jaccard', self.fill_secondary)
        cs_purple = ColorScaleRule(start_type='num', start_value=0, start_color='FFFFFF',
                                   end_type='num', end_value=1, end_color='C084FC')
        ws.conditional_formatting.add(f"C15:J22", cs_purple)

    def build_distances_sheet(self):
        ws = self.wb.create_sheet(title="Geometric Distances")
        ws.sheet_properties.tabColor = "047857"
        ws.views.sheetView[0].showGridLines = True

        r = 2
        for title_str, key in [
            ("1. EUCLIDEAN DISTANCE MATRIX", 'euclidean'),
            ("2. MANHATTAN (CITYBLOCK) DISTANCE MATRIX", 'manhattan'),
            ("3. CHEBYSHEV DISTANCE MATRIX", 'chebyshev'),
            ("4. MINKOWSKI DISTANCE MATRIX (p=3)", 'minkowski')
        ]:
            ws.cell(row=r, column=2, value=title_str).font = self.style_sec
            ws.cell(row=r + 1, column=2, value="Docs")
            for c_i, d_id in enumerate(self.engine.doc_ids, start=3):
                ws.cell(row=r + 1, column=c_i, value=d_id)

            r_idx = r + 2
            df_mat = self.engine.matrices[key]
            for d1 in self.engine.doc_ids:
                ws.cell(row=r_idx, column=2, value=d1)
                for c_i, d2 in enumerate(self.engine.doc_ids, start=3):
                    ws.cell(row=r_idx, column=c_i, value=df_mat.loc[d1, d2])
                r_idx += 1

            self._apply_format(ws, r + 1, 2, r_idx - 1, len(self.engine.doc_ids) + 2, self.fill_tertiary, num_fmt="0.0000")
            r = r_idx + 2

    def build_dtm_sheet(self):
        ws = self.wb.create_sheet(title="Term Frequency (DTM)")
        ws.sheet_properties.tabColor = "64748B"
        ws.views.sheetView[0].showGridLines = True

        ws.cell(row=2, column=2, value="DOCUMENT-TERM FREQUENCY MATRIX").font = self.style_sec
        ws.cell(row=3, column=2, value="Doc / Vocabulary")
        for c_i, word in enumerate(self.engine.feature_names, start=3):
            ws.cell(row=3, column=c_i, value=word)

        r_idx = 4
        for d_id in self.engine.doc_ids:
            ws.cell(row=r_idx, column=2, value=d_id)
            for c_i, word in enumerate(self.engine.feature_names, start=3):
                ws.cell(row=r_idx, column=c_i, value=int(self.engine.df_dtm.loc[d_id, word]))
            r_idx += 1

        self._apply_format(ws, 3, 2, r_idx - 1, len(self.engine.feature_names) + 2, self.fill_primary, num_fmt="0")

    def auto_fit_columns(self):
        for ws in self.wb.worksheets:
            for col in ws.columns:
                col_letter = get_column_letter(col[0].column)
                max_len = 0
                for cell in col:
                    if cell.value and cell.row > 3:
                        max_len = max(max_len, len(str(cell.value)))
                ws.column_dimensions[col_letter].width = max(max_len + 4, 11)

    def export(self, filename="Text_Similarity_Analytics_Dashboard.xlsx"):
        self.build_overview_sheet()
        self.build_matrices_sheet()
        self.build_distances_sheet()
        self.build_dtm_sheet()
        self.auto_fit_columns()
        self.wb.save(filename)
        return filename


# ---------------------------------------------------------
# 3. MAIN EXECUTION
# ---------------------------------------------------------
if __name__ == "__main__":
    corpus_data = {
        "D1": "penerapan algoritma machine learning untuk klasifikasi data teks",
        "D2": "pemrosesan bahasa alami menggunakan kecerdasan buatan dan machine learning",
        "D3": "pengembangan sistem rekomendasi berbasis cosine similarity dan data mining",
        "D4": "analisis sentimen media sosial menggunakan algoritma machine learning",
        "D5": "penerapan kecerdasan buatan dan machine learning pada pemrosesan bahasa alami",
        "D6": "arsitektur perangkat lunak dan rekayasa sistem basis data",
        "D7": "analisis performa algoritma pencarian pada struktur data graf",
        "D8": "pengukuran kemiripan dokumen teks menggunakan cosine similarity dan jaccard"
    }

    # Jalankan Analisis Engine
    engine = DocumentSimilarityEngine(corpus_data)
    engine.compute_similarity_and_distances()
    image_path = engine.generate_heatmap("heatmap_similarity_matrix.png")

    # Jalankan Exporter Dashboard Excel
    exporter = ExcelDashboardExporter(engine, image_path)
    excel_path = exporter.export("Text_Similarity_Analytics_Dashboard.xlsx")

    print(f"Sukses! Gambar heatmap disimpan ke: {image_path}")
    print(f"Sukses! Dashboard Excel disimpan ke: {excel_path}")
